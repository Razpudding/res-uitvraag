"""Layout definition of the Projectoverzicht (uitvraag) workbook.

Originally taken from "Projectoverzicht template najaar 2026.xlsx"; this file is now
the source of truth. Edit COLUMNS / GROUPS here to change the generated files.
"""

from dataclasses import dataclass

from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.table import TableStyleInfo

SHEET_TITLE = "Data Veld+Wind"
GROUP_ROW, HEADER_ROW, UNITS_ROW, FIRST_DATA_ROW = 1, 3, 4, 5

INT, DATE, YEAR, PCT, NUM = "0", "[$-413]mmm-yy", "0", "0%", "0.00"
GENERAL = "General"


@dataclass
class Col:
    header: str
    source: str | None      # monitor header; None = exported empty; "=fase_id"/"=fase" derived
    width: float
    fmt: str = GENERAL
    units: str | None = None
    hidden: bool = False


COLUMNS = [
    Col("Nr.", "ID", 6.2),
    Col("Projectnaam", "Projectnaam", 36.4),
    Col("Gemeente", "Gemeente", 13.6, hidden=True),
    Col("Plaats", "Plaats", 11.8),
    Col("Adres", "Adres", 13.8),
    Col("Postcode", "Postcode", 12.4),
    Col("X_coordinaat", "X_coordinaat", 11.2, hidden=True),
    Col("Y_coordinaat", "Y_coordinaat", 11.4, hidden=True),
    Col("Type", "Type", 13.4),
    Col("STATUS RES", "STATUS RES", 22.2),
    Col("Fase_ID", "=fase_id", 6.8),
    Col("Fase", "=fase", 18.2),
    Col("F1", "F1", 9.2, DATE),
    Col("F2", "F2", 9.2, DATE),
    Col("F3", "F3", 9.2, DATE),
    Col("F4", "F4", 9.2, DATE),
    Col("F5", "F5", 9.2, DATE),
    Col("F6", "F6", 9.2, DATE),
    Col("Verwachte realisatie", "verwachte_realisatie", 13.8, YEAR),
    Col("Slaagkans", "Slaagkans", 13.0, PCT, "%"),
    Col("Vermogen", "Vermogen", 21.2, NUM, "MW"),
    Col("Productie", "Productie", 20.8, NUM, "GWh_per_jaar"),
    Col("Prognose_GWh", "Prognose_GWh", 14.8, NUM, "GWh_per_jaar"),
    Col("Energie", "Energie", 15.4, NUM, "TJ_per_jaar"),
    Col("Prognose_productie", "Prognose_productie", 21.4, NUM, "TJ_per_jaar"),
    Col("Onderstation", "Onderstation", 26.8),
    Col("Ontwikkelaar", "Ontwikkelaar", 16.4),
    Col("Type zonnepark", "Type zonnepark", 13.8, units="multifunctioneel of monofunctioneel"),
    Col("Terreinomvang", "Terreinomvang", 26.2, units="(Hectare)"),
    Col("Oppervlak zonnepark", "Oppervlak zonnepark", 14.4, units="Hectare"),
    Col("Opstelling", "Opstelling", 20.6),
    Col("Orientatie", "Orientatie", 26.2),
    Col("Aantal_panelen", "Aantal_panelen", 14.4),
    Col("Aantal turbines", "Aantal turbines", 13.8),
    Col("Ashoogte", "Ashoogte", 16.6),
    Col("Tiphoogte", "Tiphoogte", 13.4),
    Col("Dakoppervlakte", "Dakoppervlakte", 13.0),
    Col("Besparing_CO2 (kton)", "Besparing_CO2 (kton)", 13.0, NUM, "kton_per_jaar"),
    Col("Elektriciteit", "Elektriciteit", 13.0, INT, "aantal_huishoudens"),
    Col("SDE", "SDE", 13.0),
    Col("Omgevingsvergunning", "Omgevingsvergunning", 13.0),
    Col("Cooperatie", "Cooperatie", 13.0),
    Col("Lokaal eigendom", "Lokaal eigendom", 13.0),
    Col("Vorm lokaal eigendom", "Vorm lokaal eigendom", 13.0),
    Col("Toelichting_LE", "Toelichting", 13.0),
    Col("toelichtin_algemeen", "Toelichting algemeen", 13.0),
    Col("Netcongestie", "Netcongestie", 13.0),
    Col("Kolom2", None, 13.0),
    Col("Kolom3", None, 13.0),
    Col("Kolom4", None, 13.0),
    Col("Kolom5", None, 13.0),
]

# Merged group labels in row 1: (label, first column header, last column header).
GROUPS = [
    ("Locatie data", "Gemeente", "Y_coordinaat"),
    ("Status", "Type", "F6"),
    ("Productie", "Verwachte realisatie", "Prognose_GWh"),
]
NOTE = ("B1", "Wijzigingen graag markeren in ROOD")
FREEZE_AT = "STATUS RES"     # columns left of this one, and rows above the data, stay visible

FONT = Font(name="Verdana", size=8)
HEADER_FONT = Font(name="Verdana", size=8, bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
NOTE_FONT = Font(name="Arial", size=10, color="3F3F76")
NOTE_FILL = PatternFill("solid", fgColor="FFCC99")
THIN = Side(style="thin", color="7F7F7F")
NOTE_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TABLE_STYLE = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
HEADER_HEIGHT, ROW_HEIGHT, ZOOM = 46.5, 14.0, 150
