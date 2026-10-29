"""Generate a Projectoverzicht (uitvraag) workbook per municipality from the RES monitor.

The monitor is opened read-only (cached formula results); it is never saved.
The layout of the generated file is defined in layout.py.

Usage:
    python generate.py --monitor <monitor.xlsx> --gemeente Hardenberg
    python generate.py --monitor <monitor.xlsx> --all
"""

import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from layout import (
    COLUMNS, FIRST_DATA_ROW, FONT, FREEZE_AT, GROUP_ROW, GROUPS, HEADER_FILL, HEADER_FONT,
    HEADER_HEIGHT, HEADER_ROW, NOTE, NOTE_BORDER, NOTE_FILL, NOTE_FONT, ROW_HEIGHT,
    SHEET_TITLE, TABLE_STYLE, UNITS_ROW, ZOOM,
)

# --- Monitor (source) ----------------------------------------------------------
MONITOR_SHEET = "Data Veld+Wind (Q2'26)"
MONITOR_HEADER_ROW = 3
MONITOR_FIRST_DATA_ROW = 4
# The monitor's general remarks column has no header; it sits right after "Toelichting".
UNNAMED_MONITOR_COLUMN = ("Toelichting algemeen", "Toelichting")


# --- Reading the monitor -------------------------------------------------------
def norm(h):
    return str(h).strip() if h is not None else None


def monitor_columns(ws):
    cols = {norm(c.value): c.column for c in ws[MONITOR_HEADER_ROW] if norm(c.value)}
    name, anchor = UNNAMED_MONITOR_COLUMN
    cols[name] = cols[anchor] + 1
    return cols


def read_projects(monitor_path):
    wb = openpyxl.load_workbook(monitor_path, data_only=True)
    ws = wb[MONITOR_SHEET]
    cols = monitor_columns(ws)
    missing = {c.source for c in COLUMNS if c.source and not c.source.startswith("=")} - set(cols)
    if missing:
        raise SystemExit(f"Monitor is missing columns: {sorted(missing)}")
    projects = []
    for r in range(MONITOR_FIRST_DATA_ROW, ws.max_row + 1):
        pid = ws.cell(r, cols["ID"]).value
        status = ws.cell(r, cols["STATUS RES"]).value
        if pid in (None, "") or status is None:
            continue
        if str(status).strip().startswith("0"):   # 0 Project gestopt
            continue
        projects.append({h: ws.cell(r, c).value for h, c in cols.items()})
    return projects


def build_row(project):
    fase_id, _, fase = str(project["STATUS RES"]).strip().partition(" ")
    derived = {"=fase_id": fase_id, "=fase": fase.strip()}
    row = []
    for col in COLUMNS:
        if col.source is None:
            row.append(None)
        elif col.source in derived:
            row.append(derived[col.source])
        else:
            v = project[col.source]
            row.append(v.strip() if isinstance(v, str) else v)
    return row


# --- Writing the uitvraag workbook --------------------------------------------
def letter(header):
    return get_column_letter([c.header for c in COLUMNS].index(header) + 1)


def write_workbook(out_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ncols = len(COLUMNS)

    ws[NOTE[0]] = NOTE[1]
    ws[NOTE[0]].font, ws[NOTE[0]].fill, ws[NOTE[0]].border = NOTE_FONT, NOTE_FILL, NOTE_BORDER
    for label, first, last in GROUPS:
        a, b = letter(first), letter(last)
        ws[f"{a}{GROUP_ROW}"] = label
        ws[f"{a}{GROUP_ROW}"].font = FONT
        ws[f"{a}{GROUP_ROW}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"{a}{GROUP_ROW}:{b}{GROUP_ROW}")

    for i, col in enumerate(COLUMNS, start=1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = col.width
        ws.column_dimensions[L].hidden = col.hidden
        for r, value in ((HEADER_ROW, col.header), (UNITS_ROW, col.units)):
            cell = ws.cell(r, i, value)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=(r == HEADER_ROW))

    for n, values in enumerate(rows):
        r = FIRST_DATA_ROW + n
        for i, (col, v) in enumerate(zip(COLUMNS, values), start=1):
            cell = ws.cell(r, i, v)
            cell.font = FONT
            cell.number_format = col.fmt
            cell.alignment = Alignment(vertical="center",
                                       horizontal="right" if isinstance(v, (int, float)) else "left")

    for r in range(GROUP_ROW, FIRST_DATA_ROW + len(rows)):
        ws.row_dimensions[r].height = ROW_HEIGHT
    ws.row_dimensions[HEADER_ROW].height = HEADER_HEIGHT

    last_row = FIRST_DATA_ROW + max(len(rows), 1) - 1
    table = Table(displayName="Projecten", ref=f"A{HEADER_ROW}:{get_column_letter(ncols)}{last_row}")
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)

    ws.freeze_panes = f"{letter(FREEZE_AT)}{FIRST_DATA_ROW}"
    ws.sheet_view.zoomScale = ZOOM

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--monitor", required=True, type=Path)
    ap.add_argument("--round", default="najaar 2026")
    ap.add_argument("--out-dir", default=Path("output"), type=Path)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--gemeente")
    g.add_argument("--all", action="store_true")
    args = ap.parse_args()

    projects = read_projects(args.monitor)
    gemeenten = sorted({p["Gemeente"] for p in projects if p["Gemeente"]}) if args.all else [args.gemeente]
    for gm in gemeenten:
        rows = [build_row(p) for p in projects if p["Gemeente"] == gm]
        if not rows:
            print(f"!! {gm}: no active projects found")
            continue
        out = args.out_dir / f"Projectoverzicht {args.round} - {gm}.xlsx"
        write_workbook(out, rows)
        print(f"{gm}: {len(rows)} projects -> {out}")


if __name__ == "__main__":
    main()
